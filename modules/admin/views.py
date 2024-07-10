import csv
import io
import json
from datetime import datetime, time
from io import StringIO, BytesIO
from typing import List, Dict, Any, Union
from typing import Optional

import pandas as pd
from flask import Blueprint, flash, redirect, url_for
from flask import request, send_file
from flask.typing import ResponseValue
from flask_admin import Admin, AdminIndexView
from flask_admin import BaseView, expose
from flask_admin.contrib.sqla import ModelView
from flask_babel import gettext as _
from flask_login import current_user
from markupsafe import Markup
from openpyxl import Workbook
from sqlalchemy import func, Integer
from sqlalchemy import inspect, select, Table
from sqlalchemy.orm import aliased
from ydata_profiling import ProfileReport

import config
from modules.db.database import db, Base
from modules.db.models import RequestLog, Ticket, User, Goods, Category, Purchase, Review, Wishlist, Tag, \
    ProductPromotion, Discount, ShippingMethod, ReportedReview, TicketMessage
from modules.decorators import login_required_with_message, admin_required

# Create Blueprint
admin_bp = Blueprint('admin', __name__, url_prefix='/admin')


def get_table_names() -> List[str]:
    """
    Get a list of all table names in the database.

    Returns:
        List[str]: A list of table names.
    """
    inspector = inspect(db.session.bind)
    return inspector.get_table_names()


def _number_formatter(view: ModelView, context: Dict, model: Any, name: str) -> str:
    """
    Format a number to display two decimal places.

    Args:
        view (ModelView): The current view.
        context (Dict): The context dictionary.
        model (Any): The model instance.
        name (str): The name of the attribute to format.

    Returns:
        str: The formatted number as a string.
    """
    value = getattr(model, name)
    if value is not None:
        return '{:.2f}'.format(value)
    return ''


class AdminView(ModelView):
    def is_accessible(self) -> bool:
        return current_user.is_authenticated and current_user.is_admin

    def inaccessible_callback(self, name: str, **kwargs) -> ResponseValue:
        return redirect(url_for('login', next=request.url))


class MyAdminIndexView(AdminIndexView):
    @expose('/')
    @login_required_with_message()
    @admin_required()
    def index(self) -> ResponseValue:
        if not current_user.is_admin:
            flash(_('You do not have permission to access this page.'), 'danger')
            return redirect(url_for('login'))
        return self.render('admin/index.html')


class TicketView(AdminView):
    column_searchable_list: list[str] = ['user.username', 'title', 'description']
    column_filters: list[str] = ['status', 'priority']
    column_editable_list: list[str] = ['status', 'priority']
    form_excluded_columns: list[str] = ['messages']

    column_formatters = {
        _('actions'): lambda v, c, m, p: Markup(
            f'<a href="{url_for("tickets.ticket_details", ticket_id=m.id)}" class="btn btn-primary btn-sm">{_("Details")}</a>'),
    }
    column_list: list[str] = ['id', 'user.username', 'title', 'status', 'priority', 'created_at', 'actions']

    @expose('/assign/<int:ticket_id>', methods=['POST'])  # type: ignore
    def assign_ticket(self, ticket_id: int) -> ResponseValue:
        ticket: Optional[Ticket] = db.session.get(Ticket, ticket_id)
        if ticket:
            admin_id = request.form.get('admin_id', type=int)
            ticket.admin_id = admin_id
            db.session.commit()
            flash(_('Ticket assigned successfully.'), 'success')
        else:
            flash(_('Ticket not found.'), 'danger')
        return redirect(url_for('ticket.index_view'))


class TicketMessageView(AdminView):
    column_searchable_list = ['user.username', 'message']
    column_filters = ['is_admin']
    form_excluded_columns = ['user', 'ticket']


class UserView(AdminView):
    column_searchable_list = ['username', 'email']
    column_filters = ['is_admin', 'is_active']
    column_editable_list = ['is_admin', 'is_active']
    form_excluded_columns = ['password']
    # hides password column from the user view
    column_exclude_list = ['password']


class StatisticsView(BaseView):
    @expose('/', methods=['GET', 'POST'])
    @login_required_with_message()
    @admin_required()
    def index(self) -> ResponseValue:
        inspector = inspect(db.engine)
        table_names = inspector.get_table_names()

        if request.method == 'POST':
            tables = request.form.getlist('tables')
            data_percentage = int(request.form.get('data_percentage', 100))

            profile_reports: List[ProfileReport] = []
            for table_name in tables:
                if table_name in table_names:
                    table = db.metadata.tables[table_name]

                    # Create an alias for the table
                    aliased_table = aliased(table)

                    count_subq = select(func.count()).select_from(table).scalar_subquery()
                    limit = select(func.cast(count_subq * (data_percentage / 100.0), Integer)).scalar_subquery()

                    query = (
                        select(aliased_table)
                        .order_by(aliased_table.c.id.desc())
                        .limit(limit)
                    )

                    df = pd.read_sql(query, db.engine)
                    profile = ProfileReport(df, title=f"{table_name.capitalize()} Dataset")
                    profile_reports.append(profile)

            if profile_reports:
                merged_report: ProfileReport = profile_reports[0]
                for report in profile_reports[1:]:
                    merged_report.add_report(report)

                profile_html = merged_report.to_html()

                output = io.BytesIO()
                output.write(profile_html.encode('utf-8'))
                output.seek(0)
                return send_file(output, mimetype='text/html', as_attachment=True,
                                 download_name='statistics_report.html')

        return self.render('admin/statistics.html', table_names=table_names)


class ReportsView(BaseView):
    @expose('/', methods=['GET', 'POST'])  # type: ignore
    @login_required_with_message()
    @admin_required()
    def index(self) -> ResponseValue:
        inspector = inspect(db.engine)
        table_names = inspector.get_table_names()

        if request.method == 'POST':
            tables = request.form.getlist('tables')
            file_format = request.form.get('file_format')

            data = self.fetch_data(tables)

            if file_format == 'csv':
                return self.generate_csv(data)
            elif file_format == 'json':
                return self.generate_json(data)
            elif file_format == 'excel':
                return self.generate_excel(data)

        return self.render('admin/reports.html', table_names=table_names)

    def fetch_data(self, tables: List[str]) -> Dict[str, List[Dict[str, Any]]]:
        inspector = inspect(db.engine)
        all_table_names = inspector.get_table_names()
        data = {}

        for table_name in tables:
            if table_name in all_table_names:
                table = Table(table_name, Base.metadata, autoload_with=db.engine)
                aliased_table = aliased(table)
                query = select(aliased_table)
                with db.engine.connect() as connection:
                    result = connection.execute(query)
                    columns = result.keys()
                    data[table_name] = [
                        {col: self.decode_if_bytes(value) for col, value in zip(columns, row)}
                        for row in result.fetchall()
                    ]
        return data

    @staticmethod
    def decode_if_bytes(value: Any) -> Any:
        return value.decode('utf-8') if isinstance(value, bytes) else value

    def generate_csv(self, data: Dict[str, List[Dict[str, Any]]]) -> ResponseValue:
        output = StringIO()
        for table, rows in data.items():
            if rows:
                writer = csv.DictWriter(output, fieldnames=rows[0].keys())
                writer.writeheader()
                writer.writerows(rows)
                output.write('\n')
        output.seek(0)
        return send_file(
            BytesIO(output.getvalue().encode('utf-8')),
            mimetype='text/csv; charset=utf-8',
            as_attachment=True,
            download_name='data.csv'
        )

    def generate_json(self, data: Dict[str, List[Dict[str, Any]]]) -> ResponseValue:
        class CustomEncoder(json.JSONEncoder):
            def default(self, obj) -> Any:
                if isinstance(obj, bytes):
                    return obj.decode('utf-8')
                return super().default(obj)

        output = json.dumps(data, ensure_ascii=False, indent=2, cls=CustomEncoder)
        return send_file(
            BytesIO(output.encode('utf-8')),
            mimetype='application/json; charset=utf-8',
            as_attachment=True,
            download_name='data.json'
        )

    def generate_excel(self, data: Dict[str, List[Dict[str, Any]]]) -> ResponseValue:
        wb = Workbook()
        wb.remove(wb.active)  # Remove default sheet
        for table, rows in data.items():
            ws = wb.create_sheet(title=table)
            if rows:
                headers = list(rows[0].keys())
                ws.append(headers)
                for row in rows:
                    ws.append([row[header] for header in headers])
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='data.xlsx'
        )


class GoodsView(AdminView):
    column_searchable_list = ['samplename', 'description']
    column_filters = ['onSale', 'category']
    column_editable_list = ['price', 'onSale', 'onSalePrice', 'stock']
    form_excluded_columns = ['purchase_items', 'reviews', 'wishlist_items', 'cart_items', 'recently_viewed_by']
    form_widget_args = {
        'description': {'rows': 10},
    }
    form_ajax_refs = {
        'category': {
            'fields': ['name'],
            'page_size': 10
        },
        'tags': {
            'fields': ['name'],
            'page_size': 10
        },
        'related_products': {
            'fields': ['samplename'],
            'page_size': 10
        }
    }

    column_formatters = {
        'price': _number_formatter,
        'onSalePrice': _number_formatter,
    }

    def on_model_change(self, form, model, is_created) -> None:
        if not is_created and model.stock <= config.AppConfig.LOW_STOCK_THRESHOLD:
            goods_name = model.samplename
            notification_message = _(f"Low stock alert: {goods_name} is running low on stock.")

            # Flash a message to the current admin user
            flash(notification_message, "info")


class CategoryView(AdminView):
    column_searchable_list = ['name']
    form_excluded_columns = ['goods']
    form_ajax_refs = {
        'parent': {
            'fields': ['name'],
            'page_size': 10
        }
    }


class AnalyticsView(BaseView):
    @expose('/', methods=['GET', 'POST'])
    @login_required_with_message()
    @admin_required()
    def index(self) -> ResponseValue:
        if request.method == 'POST':
            start_date_str: str = request.form.get('start_date', type=str,
                                                   default=datetime.today().strftime('%Y-%m-%d'))
            end_date_str: str = request.form.get('end_date', type=str, default=datetime.today().strftime('%Y-%m-%d'))

            # Convert start_date and end_date to datetime objects
            start_date: datetime = datetime.strptime(start_date_str, '%Y-%m-%d')
            end_date: datetime = datetime.strptime(end_date_str, '%Y-%m-%d')

            # Set the time component for start_date and end_date
            start_datetime: datetime = datetime.combine(start_date, time.min)
            end_datetime: datetime = datetime.combine(end_date, time.max)

            # Retrieve request logs within the specified date range
            request_logs_query = db.session.query(RequestLog).filter(
                RequestLog.timestamp >= start_datetime,
                RequestLog.timestamp <= end_datetime
            )

            total_requests: int = request_logs_query.count()
            request_logs: List[RequestLog] = request_logs_query.limit(10).all()

            # Calculate analytics metrics for all logs in the selected range
            average_execution_time: Union[int, str] = 0
            status_code_counts: Dict[int, int] = {}

            if total_requests > 0:
                # Calculate average execution time
                avg_execution_time: float = request_logs_query.with_entities(
                    func.avg(RequestLog.execution_time)).scalar() or 0.0
                average_execution_time = f"{round(avg_execution_time, 3)} s"

                # Calculate status code counts using grouping
                status_code_counts = dict(
                    request_logs_query.with_entities(
                        RequestLog.status_code, func.count(RequestLog.status_code)
                    )
                    .group_by(RequestLog.status_code)
                    .all()
                )

            # Render the template with the analytics data
            return self.render('admin/analytics.html', request_logs=request_logs, total_requests=total_requests,
                               average_execution_time=average_execution_time, status_code_counts=status_code_counts)

        return self.render('admin/analytics.html')


class PurchaseView(AdminView):
    column_searchable_list = ['user.username', 'date']
    column_filters = ['status']
    column_editable_list = ['status']
    form_excluded_columns = ['items']


class ReviewView(AdminView):
    column_searchable_list = ['user.username', 'goods.samplename']
    column_filters = ['moderated']
    column_editable_list = ['moderated']


class WishlistView(AdminView):
    column_searchable_list = ['user.username', 'goods.samplename']


class TagView(AdminView):
    column_searchable_list = ['name']
    form_excluded_columns = ['goods']


class ProductPromotionView(AdminView):
    form_excluded_columns = ['goods']


class DiscountView(AdminView):
    column_searchable_list = ['code']
    column_filters = ['start_date', 'end_date']
    form_excluded_columns = ['users']


class ShippingMethodView(AdminView):
    column_searchable_list = ['name']


class ReportedReviewView(AdminView):
    column_searchable_list = ['review.user.username', 'review.goods.samplename']
    column_filters = ['created_at']
    form_excluded_columns = ['user', 'review']
    column_formatters = {
        _('actions'): lambda v, c, m, p: Markup(
            f'<a href="{url_for("reviews.reported_review_detail", review_id=m.review_id)}" class="btn btn-primary btn-sm">{_("Details")}</a>'),
    }
    column_list = ['review_id', 'review.user.username', 'review.goods.samplename', 'explanation', 'created_at',
                   'actions']


############

# Create Admin instance
admin = Admin(name='Admin Panel', template_mode='bootstrap4', index_view=MyAdminIndexView())

# Add views to admin
admin.add_view(UserView(User, db.session, name='Users'))
admin.add_view(GoodsView(Goods, db.session, name='Goods'))
admin.add_view(CategoryView(Category, db.session, name='Categories'))
admin.add_view(PurchaseView(Purchase, db.session, name='Purchases'))
admin.add_view(ReviewView(Review, db.session, name='Reviews'))
admin.add_view(WishlistView(Wishlist, db.session, name='Wishlists'))
admin.add_view(TagView(Tag, db.session, name='Tags'))
admin.add_view(ProductPromotionView(ProductPromotion, db.session, name='Promotions'))
admin.add_view(DiscountView(Discount, db.session, name='Discounts'))
admin.add_view(ShippingMethodView(ShippingMethod, db.session, name='Shipping Methods'))
admin.add_view(ReportedReviewView(ReportedReview, db.session, name='Reported Reviews'))
admin.add_view(TicketView(Ticket, db.session, name='Tickets'))
admin.add_view(TicketMessageView(TicketMessage, db.session, name='Ticket Messages'))
admin.add_view(StatisticsView(name='Statistics', endpoint='statistics', category='Statistics'))
admin.add_view(ReportsView(name='Reports', endpoint='reports', category='Reports'))
admin.add_view(AnalyticsView(name='Analytics', endpoint='analytics', category='Reports'))


def init_admin(app) -> None:
    admin.init_app(app)
